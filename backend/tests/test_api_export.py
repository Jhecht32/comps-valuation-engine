"""POST /api/comps/export: the comps run as a formatted workbook with a
Comps, Statistics, Implied, and Sources sheet. Recorded Home Depot data
plus a canned proxy peer group, so the Sources sheet has candidates,
filtered names, a filing, and flags to list."""

from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.data import CompanyProfile, FixtureProvider
from app.data.fixtures import HOME_DEPOT
from app.data.proxy_peers import Confidence, FixtureProxyPeerSource, ProxyPeer, ProxyPeerGroup
from app.main import create_app
from tests.helpers import RECORDED_AT, recorded_provider

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HD_PROXY_URL = "https://www.sec.gov/Archives/edgar/data/354950/000035495026000090/hd-20260406.htm"
XOM = CompanyProfile(
    ticker="XOM", name="Exxon Mobil Corporation", sector="Energy", industry="Oil & Gas Integrated",
    market_cap=500e9, price_currency="USD", reporting_currency="USD",
)
HD_GROUP = ProxyPeerGroup(
    ticker="HD",
    company_name="HOME DEPOT, INC.",
    filing_date=date(2026, 4, 7),
    accession_number="0000354950-26-000090",
    document_url=HD_PROXY_URL,
    peers=[
        ProxyPeer(name="Lowe’s Companies, Inc.", ticker="LOW"),
        ProxyPeer(name="Tractor Supply Company", ticker="TSCO"),
        ProxyPeer(name="Exxon Mobil Corporation", ticker="XOM"),
    ],
    unmatched=["Bunnings Group"],
    confidence=Confidence.MEDIUM,
    context="Retail Peer Group",
)
REQUEST = {"target": "HD", "peers": ["LOW", "FND", "TSCO", "WSM", "ZZZZNOPE"]}
COMPS_HEADERS = [
    "Ticker", "Company", "Share price", "Diluted shares (mm)", "Equity value (mm)", "Total debt (mm)",
    "Cash & STI (mm)", "EV (mm)", "LTM revenue (mm)", "LTM EBITDA (mm)", "EBITDA margin",
    "EV / Revenue", "EV / EBITDA", "EV / EBIT", "P / E",
]
MILLIONS = "#,##0.0;(#,##0.0)"
PER_SHARE = "#,##0.00;(#,##0.00)"
MULTIPLE = '0.0"x";(0.0"x")'
PERCENT = "0.0%;(0.0%)"


@pytest.fixture
def client():
    app = create_app(
        provider=recorded_provider(XOM=XOM),
        proxy_source=FixtureProxyPeerSource({"HD": HD_GROUP}),
        today=lambda: RECORDED_AT,
    )
    with TestClient(app) as c:
        yield c


def export(client, body=REQUEST):
    r = client.post("/api/comps/export", json=body)
    assert r.status_code == 200, r.text
    return r, load_workbook(BytesIO(r.content))


def rows_of(ws):
    return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]


def test_export_is_an_xlsx_attachment_named_for_the_target_and_date(client):
    r, wb = export(client)
    assert r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == 'attachment; filename="HD_comps_2026-09-06.xlsx"'
    assert wb.sheetnames == ["Comps", "Statistics", "Implied", "Sources"]
    assert all(wb[name].freeze_panes == "A2" for name in wb.sheetnames)


def test_comps_sheet_has_the_ui_columns_and_a_distinct_populated_target_row(client):
    _, wb = export(client)
    ws = wb["Comps"]
    assert [c.value for c in ws[1]] == COMPS_HEADERS
    target = [c.value for c in ws[2]]
    assert target[0] == "HD" and target[1] == "The Home Depot, Inc."
    assert all(isinstance(v, (int, float)) for v in target[2:]), target  # every figure numeric, none NM or blank
    assert target[2] > 0 and target[3] > 100 and target[7] > target[4] > 1_000  # price; shares mm; EV > equity, in mm
    assert target[4] == pytest.approx(target[2] * target[3], rel=1e-6)  # equity value = price × diluted shares, both in mm
    assert 0 < target[10] < 1 and target[12] > 5  # margin a fraction; EV/EBITDA a multiple
    assert [ws.cell(r, 1).value for r in range(3, 7)] == ["LOW", "FND", "TSCO", "WSM"]
    assert ws.max_row == 6  # the unvaluable ticker is not a row
    # The target row is bold on the UI's shaded background; peers are plain.
    assert ws.cell(2, 1).font.bold and ws.cell(2, 2).font.bold and ws.cell(2, 1).fill.fgColor.rgb.endswith("EDF2FA")
    assert not ws.cell(3, 1).font.bold and ws.cell(3, 1).fill.fill_type is None
    # Number formats by column kind.
    assert ws.cell(2, 3).number_format == PER_SHARE
    assert ws.cell(2, 4).number_format == MILLIONS and ws.cell(3, 8).number_format == MILLIONS
    assert ws.cell(2, 11).number_format == PERCENT
    assert [ws.cell(2, c).number_format for c in range(12, 16)] == [MULTIPLE] * 4
    assert ws.column_dimensions["B"].width >= 24 and ws.column_dimensions["A"].width >= 8


def test_statistics_sheet_has_percentiles_with_n_counts(client):
    _, wb = export(client)
    ws = wb["Statistics"]
    assert [c.value for c in ws[1]] == ["Multiple", "n", "Min", "25th", "Median", "Mean", "75th", "Max"]
    assert [ws.cell(r, 1).value for r in range(2, 6)] == ["EV / Revenue", "EV / EBITDA", "EV / EBIT", "P / E"]
    ebitda = [ws.cell(3, c).value for c in range(2, 9)]
    assert ebitda[0] == 4
    assert ebitda[1] <= ebitda[2] <= ebitda[3] <= ebitda[5] <= ebitda[6]
    assert all(ws.cell(3, c).number_format == MULTIPLE for c in range(3, 9))
    assert ws.cell(3, 2).number_format == "0"


def test_implied_sheet_has_the_per_share_range_per_methodology(client):
    _, wb = export(client)
    ws = wb["Implied"]
    assert [c.value for c in ws[1]] == ["Methodology", "n", "Low (25th pct)", "Mid (peer median)", "High (75th pct)", "Note"]
    assert [ws.cell(r, 1).value for r in range(2, 6)] == ["EV / Revenue", "EV / EBITDA", "EV / EBIT", "P / E"]
    low, mid, high = (ws.cell(3, c).value for c in range(3, 6))
    assert ws.cell(3, 2).value == 4 and 0 < low <= mid <= high
    assert all(ws.cell(3, c).number_format == PER_SHARE for c in range(3, 6))
    assert ws.cell(3, 6).value is None


def test_sources_sheet_lists_dates_filing_candidates_filters_and_flags(client):
    _, wb = export(client)
    ws = wb["Sources"]
    assert [c.value for c in ws[1]] == ["Section", "Item", "Tag", "Detail"]
    rows = rows_of(ws)
    by_section = {}
    for section, *rest in rows:
        by_section.setdefault(section, []).append(tuple(rest))
    assert list(by_section) == ["Run", "Proxy statement", "Candidates", "Filtered out", "Added by hand", "Flags", "Peers not valued"]

    run = dict((item, detail) for item, _, detail in by_section["Run"])
    assert run["As of"] == "2026-09-06" and run["Mid basis"] == "median"
    assert run["HD share price as of"] == "2026-09-05" or run["HD share price as of"] <= "2026-09-06"
    assert run["HD balance sheet as of"] and run["HD LTM through"]
    assert run["Peers requested"] == "LOW, FND, TSCO, WSM, ZZZZNOPE"

    proxy = dict((item, (tag, detail)) for item, tag, detail in by_section["Proxy statement"])
    assert proxy["Filing"] == ("DEF 14A", HD_PROXY_URL)
    assert proxy["Filed"] == (None, "2026-04-07") and proxy["Confidence"] == (None, "medium")
    assert proxy["Unmatched names"] == (None, "Bunnings Group")

    candidates = by_section["Candidates"]
    assert candidates[0][0] == "LOW  Lowe's Companies, Inc." and candidates[0][1] == "proxy + screen"
    assert "Home Improvement Retail" in candidates[0][2] and "in this run" in candidates[0][2]

    filtered = {item.split("  ")[0]: (tag, detail) for item, tag, detail in by_section["Filtered out"]}
    assert filtered["TSCO"][0] == "size" and filtered["TSCO"][1].startswith("market cap $18bn against HD's $320bn")
    assert filtered["XOM"] == ("sector", "Energy sector (Oil & Gas Integrated), not Consumer Cyclical")

    assert [item for item, _, _ in by_section["Added by hand"]] == ["FND", "TSCO", "WSM", "ZZZZNOPE"]

    flags = [(item, tag) for item, tag, _ in by_section["Flags"]]
    assert ("HD", "proxy_peers_unmatched") in flags and ("HD", "proxy_peers_dropped") in flags
    assert by_section["Peers not valued"][0][:2] == ("ZZZZNOPE", "ticker_not_found")
    assert ws.column_dimensions["D"].width >= 60


def test_nm_is_written_as_text_and_a_missing_suggestion_is_noted():
    # A target with negative EBITDA has NM EV/EBITDA and EV/EBIT, in the
    # comps row and in the implied range. No profiles means the peer
    # suggestion cannot run; the Sources sheet says so instead of failing.
    losing = HOME_DEPOT.model_copy(update={
        "ticker": "NEG", "name": "Neg Inc.",
        "quarters": [q.model_copy(update={"ebit": -(q.d_and_a + 1_000_000_000.0)}) for q in HOME_DEPOT.quarters],
    }, deep=True)
    app = create_app(provider=FixtureProvider({"HD": HOME_DEPOT, "NEG": losing}), today=lambda: RECORDED_AT)
    with TestClient(app) as c:
        _, wb = export(c, {"target": "NEG", "peers": ["HD"]})
    ws = wb["Comps"]
    assert ws.cell(2, 1).value == "NEG" and ws.cell(2, 13).value == "NM" and ws.cell(2, 14).value == "NM"
    assert isinstance(ws.cell(2, 12).value, (int, float))  # EV/Revenue still a number
    assert ws.cell(2, 10).value < 0 and ws.cell(2, 10).number_format == MILLIONS  # negative EBITDA, shown in parentheses
    implied = wb["Implied"]
    assert [implied.cell(3, c).value for c in range(3, 6)] == ["NM", "NM", "NM"]
    assert implied.cell(3, 6).value  # the reason travels with it
    sources = rows_of(wb["Sources"])
    assert any(r[0] == "Run" and r[1] == "Peer suggestion" and "unavailable" in r[3] for r in sources)
    assert not any(r[0] == "Candidates" for r in sources)


def test_export_target_not_found_is_the_usual_json_error(client):
    r = client.post("/api/comps/export", json={"target": "ZZZZNOPE", "peers": ["LOW"]})
    assert r.status_code == 404 and r.json()["code"] == "ticker_not_found" and r.json()["ticker"] == "ZZZZNOPE"


def test_openapi_documents_the_export_as_a_spreadsheet(client):
    op = client.get("/openapi.json").json()["paths"]["/api/comps/export"]["post"]
    assert XLSX in op["responses"]["200"]["content"]
    assert "404" in op["responses"]
