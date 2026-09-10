"""A comps run as an Excel workbook: the UI's tables plus a Sources sheet
that carries the audit trail (dates, the proxy filing, every candidate
with its source tags, every name filtered out with its reason, every
flag), so the file stands on its own.

Numbers are written as numbers with Excel formats, never as formatted
strings, so the reader can compute on them: money in millions with
thousands separators, multiples as 0.0x, margins as 0.0%, negatives in
parentheses. NM is the text "NM"; a figure the source did not report is
left blank.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.models import CompsResult, PeerSuggestions
from app.services.peers import format_cap
from app.valuation.models import MultipleStats, MultipleValue

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MILLIONS = "#,##0.0;(#,##0.0)"
PER_SHARE = "#,##0.00;(#,##0.00)"
MULTIPLE = '0.0"x";(0.0"x")'
PERCENT = "0.0%;(0.0%)"
COUNT = "0"
NM = "NM"
MM = 1e6

MULTIPLES = [("ev_revenue", "EV / Revenue"), ("ev_ebitda", "EV / EBITDA"), ("ev_ebit", "EV / EBIT"), ("pe", "P / E")]
COMPS_HEADERS = [
    "Ticker", "Company", "Share price", "Diluted shares (mm)", "Equity value (mm)", "Total debt (mm)",
    "Cash & STI (mm)", "EV (mm)", "LTM revenue (mm)", "LTM EBITDA (mm)", "EBITDA margin",
    "EV / Revenue", "EV / EBITDA", "EV / EBIT", "P / E",
]
COMPS_FORMATS = [None, None, PER_SHARE] + [MILLIONS] * 7 + [PERCENT] + [MULTIPLE] * 4
COMPS_WIDTHS = [10, 30, 12, 18, 16, 14, 14, 14, 16, 16, 14, 13, 13, 13, 13]

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="F5F5F2")  # the UI's panel colour
_TARGET_FILL = PatternFill("solid", fgColor="EDF2FA")  # the UI's shaded target row
_WRAP = Alignment(wrap_text=True, vertical="top")


def export_filename(ticker: str, as_of: date) -> str:
    return f"{ticker}_comps_{as_of.isoformat()}.xlsx"


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    return ws


def _mm(raw: float | None) -> float | None:
    return None if raw is None else raw / MM


def _multiple(m: MultipleValue) -> float | str:
    return NM if m.is_nm or m.value is None else m.value


def _stat(v: float | None) -> float | str:
    return NM if v is None else v


def _margin(ebitda: float | None, revenue: float | None) -> float | str | None:
    if ebitda is None or revenue is None:
        return None
    return NM if revenue <= 0 else ebitda / revenue


def _write(ws: Worksheet, values: list, formats: list[str | None], *, bold: bool = False, fill: PatternFill | None = None) -> int:
    """Append one row, formatting numeric cells; returns the row number."""
    ws.append(values)
    row = ws.max_row
    for col, fmt in enumerate(formats, start=1):
        cell = ws.cell(row, col)
        if fmt and isinstance(cell.value, (int, float)):
            cell.number_format = fmt
        if bold:
            cell.font = Font(bold=True)
        if fill is not None:
            cell.fill = fill
    return row


def _comps_sheet(wb: Workbook, result: CompsResult) -> None:
    ws = _sheet(wb, "Comps", COMPS_HEADERS, COMPS_WIDTHS)
    for i, c in enumerate([result.target, *result.peers]):
        b, ltm, m = c.bridge, c.ltm, c.multiples
        row = _write(
            ws,
            [
                c.ticker, c.name, c.share_price, _mm(b.diluted_shares), _mm(b.equity_value), _mm(b.total_debt),
                _mm(b.cash_and_equivalents + b.short_term_investments), _mm(b.enterprise_value),
                _mm(ltm.revenue), _mm(ltm.ebitda), _margin(ltm.ebitda, ltm.revenue),
                _multiple(m.ev_revenue), _multiple(m.ev_ebitda), _multiple(m.ev_ebit), _multiple(m.pe),
            ],
            COMPS_FORMATS,
            bold=i == 0,
            fill=_TARGET_FILL if i == 0 else None,
        )
        if b.shares_are_basic_fallback:
            ws.cell(row, 4).comment = Comment("Basic shares; diluted count unavailable", "comps")


def _statistics_sheet(wb: Workbook, result: CompsResult) -> None:
    ws = _sheet(wb, "Statistics", ["Multiple", "n", "Min", "25th", "Median", "Mean", "75th", "Max"], [16, 6, 10, 10, 10, 10, 10, 10])
    for key, label in MULTIPLES:
        s: MultipleStats = getattr(result.statistics, key)
        _write(ws, [label, s.n, _stat(s.min), _stat(s.p25), _stat(s.median), _stat(s.mean), _stat(s.p75), _stat(s.max)], [None, COUNT] + [MULTIPLE] * 6)


def _implied_sheet(wb: Workbook, result: CompsResult) -> None:
    basis = result.implied.mid_basis.value
    ws = _sheet(wb, "Implied", ["Methodology", "n", "Low (25th pct)", f"Mid (peer {basis})", "High (75th pct)", "Note"], [16, 6, 16, 18, 16, 60])
    for key, label in MULTIPLES:
        rng = getattr(result.implied, key)
        if rng.is_nm:
            _write(ws, [label, rng.n, NM, NM, NM, rng.nm_reason or "Not meaningful"], [None, COUNT])
        else:
            _write(ws, [label, rng.n, rng.low, rng.mid, rng.high, None], [None, COUNT, PER_SHARE, PER_SHARE, PER_SHARE])


def _sources_sheet(
    wb: Workbook, result: CompsResult, suggestions: PeerSuggestions | None, requested: list[str], suggestion_error: str | None
) -> None:
    ws = _sheet(wb, "Sources", ["Section", "Item", "Tag", "Detail"], [18, 36, 18, 100])
    t = result.target

    def add(section: str, item: str | None, tag: str | None, detail: str | None) -> None:
        ws.append([section, item, tag, detail])
        ws.cell(ws.max_row, 4).alignment = _WRAP

    add("Run", "As of", None, result.as_of.isoformat())
    add("Run", "Mid basis", None, result.implied.mid_basis.value)
    add("Run", f"{t.ticker} share price as of", None, t.price_as_of.isoformat())
    add("Run", f"{t.ticker} balance sheet as of", None, t.balance_sheet_as_of.isoformat())
    add("Run", f"{t.ticker} LTM through", None, t.ltm.as_of.isoformat())
    add("Run", "Peers requested", None, ", ".join(requested))
    add("Run", "Peers valued", None, ", ".join(p.ticker for p in result.peers) or "none")
    if suggestions is None:
        add("Run", "Peer suggestion", None, f"unavailable: {suggestion_error or 'no peer suggestion for this target'}")

    if suggestions is not None:
        proxy = suggestions.proxy
        if proxy is None:
            add("Proxy statement", "Filing", None, "No proxy-statement peer group was read; see Flags.")
        else:
            add("Proxy statement", "Filing", "DEF 14A", proxy.document_url)
            add("Proxy statement", "Filed", None, proxy.filing_date.isoformat())
            add("Proxy statement", "Confidence", None, proxy.confidence.value)
            add("Proxy statement", "Found under", None, proxy.context)
            add("Proxy statement", "Names in the filing", None, ", ".join(p.ticker for p in proxy.peers) or "none mapped")
            if proxy.unmatched:
                add("Proxy statement", "Unmatched names", None, ", ".join(proxy.unmatched))

        used = set(requested)
        for peer in suggestions.peers:
            detail = [peer.industry or peer.sector or "no classification"]
            if peer.market_cap is not None:
                detail.append("market cap " + format_cap(peer.market_cap, peer.price_currency))
            if peer.ebitda_margin is not None:
                detail.append(f"EBITDA margin {peer.ebitda_margin * 100:.1f}%")
            detail.append("in this run" if peer.ticker in used else "removed by hand")
            add("Candidates", f"{peer.ticker}  {peer.name or ''}".rstrip(), " + ".join(s.value for s in peer.sources), " · ".join(detail))

        dropped = {d.ticker: d for d in suggestions.proxy_dropped}
        for d in suggestions.proxy_dropped:
            add("Filtered out", f"{d.ticker}  {d.name}", d.rule.value, d.reason)

        candidates = {p.ticker for p in suggestions.peers}
        for ticker in requested:
            if ticker in candidates:
                continue
            note = "not among the suggested candidates"
            if ticker in dropped:
                note += f"; filtered out of the proxy group ({dropped[ticker].rule.value}) and added back"
            add("Added by hand", ticker, None, note)

    flags = [(t.ticker, f) for f in (suggestions.flags if suggestions is not None else [])]
    for c in [t, *result.peers]:
        flags.extend((c.ticker, f) for f in c.flags)
    for ticker, f in flags:
        add("Flags", ticker, f.code.value, f.message)
    if not flags:
        add("Flags", "—", None, "No data-quality flags.")

    for ticker, err in result.errors.items():
        add("Peers not valued", ticker, err.code.value, err.message)
    if not result.errors:
        add("Peers not valued", "—", None, "Every requested peer was valued.")


def build_workbook(
    result: CompsResult,
    suggestions: PeerSuggestions | None,
    *,
    requested_peers: list[str],
    suggestion_error: str | None = None,
) -> bytes:
    """The four sheets as .xlsx bytes. suggestions is the target's peer
    suggestion for the Sources sheet, or None with suggestion_error
    saying why it could not be made."""
    wb = Workbook()
    wb.remove(wb.active)
    _comps_sheet(wb, result)
    _statistics_sheet(wb, result)
    _implied_sheet(wb, result)
    _sources_sheet(wb, result, suggestions, requested_peers, suggestion_error)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
